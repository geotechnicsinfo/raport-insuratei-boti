#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Construiește data.json și pozele optimizate pentru aplicația de raportare
a investigațiilor de teren – Însurăței.

Surse:
  - fise foraj Insuratei.xlsx  (foaia "Fise": stratificație, apă, dată, GPS)
  - foraje facute.xlsx          (lista investigațiilor executate)
  - OSTSG – ... .kml            (punctele proiectate: Fr, D, FT, Fp, CPTu, Rt, DH, Fb)
  - foldere "Foraj Nxx" / "Dezvelire Nxx" cu poze (WhatsApp) + fișa manuscrisă
  - meteo.json                  (Open-Meteo archive, opțional)

Rulare:
  python3 scripts/build_data.py --fise "<cale xlsx>" --facute "<cale xlsx>" \
      --kml "<cale kml>" --poze "<folder poze>" --out "<folder app>"
"""
import argparse, glob, html, json, os, re, sys, datetime
import openpyxl
from PIL import Image, ImageOps

MAX_W = 1200   # px, latura lungă a pozelor
THUMB_W = 420
QUALITY = 68


def dms_to_dec(s):
    """'N 44°51'43.5"' -> 44.862083"""
    if not s:
        return None
    m = re.search(r'([NSEW])\s*(\d+)[°\s]+(\d+)[\'\s]+([\d.]+)', str(s))
    if not m:
        return None
    h, d, mi, se = m.group(1), float(m.group(2)), float(m.group(3)), float(m.group(4))
    v = d + mi / 60 + se / 3600
    return round(-v if h in 'SW' else v, 6)


def norm_date(s):
    """'29.09.2026' -> '2026-09-29'"""
    if not s:
        return None
    if isinstance(s, (datetime.date, datetime.datetime)):
        return s.strftime('%Y-%m-%d')
    m = re.match(r'(\d{1,2})\.(\d{1,2})\.(\d{4})', str(s).strip())
    if not m:
        return None
    return f'{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}'


def parse_fise(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['Fise']
    inv = {}
    for r in ws.iter_rows(min_row=1, values_only=True):
        name = r[0]
        if not name or name in ('foraj', 'Amplasament:'):
            continue
        name = str(name).strip()
        d = inv.setdefault(name, {'id': name, 'straturi': [], 'apa': None, 'data': None,
                                  'gps_dms': [], 'note_apa': []})
        de, la, desc = r[1], r[2], r[4]
        if desc:
            d['straturi'].append({'de': de, 'la': la, 'descriere': str(desc).strip()})
        if r[5]:
            if d['apa'] is None:
                d['apa'] = str(r[5]).strip()
            else:
                d['note_apa'].append(str(r[5]).strip())
        if r[6]:
            d['data'] = norm_date(r[6])
        if r[7]:
            d['gps_dms'].append(str(r[7]).strip())
    for d in inv.values():
        lat = lon = None
        for g in d['gps_dms']:
            v = dms_to_dec(g)
            if g.strip().startswith('N'):
                lat = v
            elif g.strip().startswith('E'):
                lon = v
        d['lat'], d['lon'] = lat, lon
        # adâncime finală
        la_vals = [s['la'] for s in d['straturi'] if isinstance(s['la'], (int, float))]
        de_vals = [s['de'] for s in d['straturi'] if isinstance(s['de'], (int, float))]
        d['adancime_fisa'] = max(la_vals + de_vals) if (la_vals or de_vals) else None
        # NH / IN
        allw = ' '.join([d['apa'] or ''] + d['note_apa'])
        m = re.search(r'NH\s*=\s*([\d.,]+)', allw)
        d['NH'] = float(m.group(1).replace(',', '.')) if m else None
        m = re.search(r'IN\s*=\s*([\d.,]+)', allw)
        d['IN'] = float(m.group(1).replace(',', '.')) if m else None
        # probe (ștuț / macroporic) deduse din descrieri
        probe = []
        for s in d['straturi']:
            txt = s['descriere'].lower()
            if re.search(r'ștuț|stut|stuț|\(stat|macroporic', txt):
                m = re.search(r'la\s*([\d,\.]+)\s*m', txt)
                ad = float(m.group(1).replace(',', '.')) if m else s['de']
                probe.append({'adancime': ad, 'tip': 'ștuț (probă netulburată)',
                              'observatii': 'strat macroporic' if 'macroporic' in txt else ''})
        d['probe'] = probe
    return inv


def parse_facute(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    out = {}
    for r in ws.iter_rows(min_row=3, values_only=True):
        for off in (0, 6):
            try:
                tip, nr, data, ad = r[off + 1], r[off + 2], r[off + 3], r[off + 4]
            except IndexError:
                continue
            if tip in ('D', 'F') and nr:
                out[f'{tip}{int(nr)}'] = {'data': norm_date(data), 'adancime': ad}
    return out


def parse_kml(path):
    s = open(path, encoding='utf-8').read()
    pts = []
    for p in re.findall(r'<Placemark[^>]*>(.*?)</Placemark>', s, re.S):
        n = re.search(r'<name>(.*?)</name>', p, re.S)
        c = re.search(r'<Point>.*?<coordinates>(.*?)</coordinates>', p, re.S)
        if not (n and c):
            continue
        d = re.search(r'<description>(.*?)</description>', p, re.S)
        desc = html.unescape(d.group(1)) if d else ''
        desc = desc.replace('<![CDATA[', '').replace(']]>', '')
        desc = re.sub(r'<br\s*/?>', '\n', desc)
        desc = re.sub(r'<.*?>', '', desc)
        lon, lat = [float(x) for x in c.group(1).strip().split(',')[:2]]
        f = lambda k: (re.search(k + r':\s*(.*)', desc) or [None, None])[1]
        st = re.search(r'E\s*=\s*([\d.]+)\s*m;\s*N\s*=\s*([\d.]+)', desc)
        pts.append({'id': n.group(1).strip(), 'lat': lat, 'lon': lon,
                    'tip': (f('Tip') or '').strip(),
                    'adancime': (f(r'Adâncime(?: foraj)?') or '').strip(),
                    'obs': (f('Observații') or '').strip(),
                    'stereo70': {'E': float(st.group(1)), 'N': float(st.group(2))} if st else None})
    return pts


def process_photos(folder, out_dir, rel):
    """Redimensionează pozele; returnează (fisa, [poze])"""
    os.makedirs(out_dir, exist_ok=True)
    files = sorted(f for f in os.listdir(folder) if re.search(r'\.jpe?g$|\.png$', f, re.I))
    fisa, pics = None, []
    n = 0
    for f in files:
        src = os.path.join(folder, f)
        try:
            im = ImageOps.exif_transpose(Image.open(src)).convert('RGB')
        except Exception as e:
            print('  ! poza ignorată', src, e)
            continue
        is_fisa = 'fisa' in f.lower() or bool(re.match(r'^(D|F|Fr)[\s_-]*\d+\.(jpe?g|png)$', f, re.I))
        name = 'fisa' if is_fisa else f'{n + 1:02d}'
        for w, suffix in ((MAX_W, ''), (THUMB_W, 't_')):
            img = im.copy()
            img.thumbnail((w, w))
            dst = os.path.join(out_dir, f'{suffix}{name}.jpg')
            if os.path.exists(dst) and os.path.getmtime(dst) >= os.path.getmtime(src):
                continue   # deja procesată
            img.save(dst, 'JPEG', quality=QUALITY, optimize=True, progressive=True)
        entry = {'src': f'{rel}/{name}.jpg', 'thumb': f'{rel}/t_{name}.jpg',
                 'original': f, 'w': im.width, 'h': im.height}
        m = re.search(r'(\d{4}-\d{2}-\d{2}) at (\d{2})\.(\d{2})', f)
        if m:
            entry['timestamp'] = f'{m.group(1)} {m.group(2)}:{m.group(3)}'
        if is_fisa:
            fisa = entry
        else:
            pics.append(entry)
            n += 1
    return fisa, pics


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--fise', required=True)
    ap.add_argument('--facute', required=True)
    ap.add_argument('--kml', required=True)
    ap.add_argument('--poze', required=True)
    ap.add_argument('--meteo', default=None)
    ap.add_argument('--out', required=True)
    a = ap.parse_args()

    fise = parse_fise(a.fise)
    facute = parse_facute(a.facute)
    kml = parse_kml(a.kml)
    kml_by_id = {p['id']: p for p in kml}
    print(f'fise: {len(fise)}  facute: {len(facute)}  kml: {len(kml)}')

    ids = sorted(set(fise) | set(facute), key=lambda x: (x[0], int(re.sub(r'\D', '', x) or 0)))
    investigatii = []
    for iid in ids:
        tip = 'dezvelire' if iid.startswith('D') else 'foraj'
        nr = int(re.sub(r'\D', '', iid))
        f = fise.get(iid, {})
        fc = facute.get(iid, {})
        planned = kml_by_id.get(iid if tip == 'dezvelire' else f'Fr{nr}')
        folder_name = ('Dezvelire' if tip == 'dezvelire' else 'Foraj') + f' N{nr}'
        folder = os.path.join(a.poze, folder_name)
        fisa_img, pics = (None, [])
        pid = planned['id'] if planned else iid
        if os.path.isdir(folder):
            print('poze', folder_name)
            fisa_img, pics = process_photos(folder, os.path.join(a.out, 'photos', pid), f'photos/{pid}')
        data = f.get('data') or fc.get('data')
        nota = None
        # corecție: fișa F92 este datată 23.10.2026 (eroare de scriere), lista de execuție indică 23.09.2026
        if f.get('data') and fc.get('data') and f['data'] != fc['data']:
            pd = sorted({x['timestamp'][:10] for x in pics if x.get('timestamp')})
            nota = (f"Data din fișă ({f['data']}) diferă de lista de execuție ({fc['data']}); s-a reținut data din lista de execuție."
                    + (f" Fotografiile au fost transmise pe {', '.join(pd)}." if pd else ''))
            data = fc['data']
        lat, lon = f.get('lat'), f.get('lon')
        coord_sursa = 'GPS teren'
        if (lat is None or lon is None) and planned:
            lat, lon, coord_sursa = planned['lat'], planned['lon'], 'punct proiectat (fără GPS în fișă)'
        rt = None
        if planned:
            m = re.search(r'\((Rt\d+)\)', planned['obs'])
            if m:
                rt = m.group(1)
        # denumirea oficială = numele punctului din KML (Fr67 / D67); fallback pe codul din fișă
        kid = planned['id'] if planned else iid
        ktip = planned['tip'] if planned else ('Sondaje deschise – dezveliri drumuri (D)' if tip == 'dezvelire' else 'Foraje drumuri (Fr)')
        investigatii.append({
            'id': kid, 'id_fisa': iid, 'tip': tip, 'nr': nr, 'categorie_kml': ktip,
            'eticheta': ('Sondaj deschis (dezvelire) ' if tip == 'dezvelire' else 'Foraj drum ') + kid,
            'data': data, 'nota_data': nota,
            'adancime': fc.get('adancime') or f.get('adancime_fisa'),
            'apa': f.get('apa'), 'NH': f.get('NH'), 'IN': f.get('IN'),
            'gps_dms': f.get('gps_dms', []), 'lat': lat, 'lon': lon, 'coord_sursa': coord_sursa,
            'proiectat': planned and {'id': planned['id'], 'lat': planned['lat'], 'lon': planned['lon'],
                                      'stereo70': planned['stereo70'], 'obs': planned['obs'],
                                      'adancime': planned['adancime']},
            'rt': rt,
            'straturi': f.get('straturi', []),
            'probe': f.get('probe', []),
            'fisa_foto': fisa_img, 'poze': pics,
        })

    meteo = {}
    if a.meteo and os.path.exists(a.meteo):
        mj = json.load(open(a.meteo, encoding='utf-8'))
        d = mj['daily']
        for i, t in enumerate(d['time']):
            meteo[t] = {k: d[k][i] for k in d if k != 'time'}
        meteo['_meta'] = {'sursa': 'Open-Meteo (ERA5 / archive API)', 'lat': mj.get('latitude'),
                          'lon': mj.get('longitude'), 'elevatie_m': mj.get('elevation')}

    # stadiu execuție pe categorii (din KML)
    from collections import Counter
    for p in kml:  # piezometrele sunt descrise individual în KML -> o singură categorie
        if p['tip'].lower().startswith('foraj fundație turbină'):
            p['tip'] = 'Foraje echipate piezometric'
    plan_cnt = Counter(p['tip'] for p in kml if p['tip'])
    exec_ids = {i['id'] for i in investigatii} | {i['rt'] for i in investigatii if i.get('rt')}
    stadiu = []
    for tipk, total in sorted(plan_cnt.items(), key=lambda x: -x[1]):
        ids = [p['id'] for p in kml if p['tip'] == tipk]
        done = [x for x in ids if x in exec_ids]
        stadiu.append({'categorie': tipk, 'proiectate': total, 'executate': len(done),
                       'neexecutate': [x for x in ids if x not in exec_ids]})
    out = {
        'stadiu': stadiu,
        'proiect': {
            'titlu': 'Investigații geotehnice Faza I – UAT Însurăței',
            'beneficiar_lucrare': 'Parc eolian – UAT Însurăței, jud. Brăila',
            'amplasament': 'Însurăței, județul Brăila',
            'executant': 'SC BBCGEOTECHNIC SRL',
            'tip_lucrari': 'Foraje drumuri de acces/exploatare (Fr) și sondaje deschise – dezveliri drumuri (D), adâncime 4,00 m',
            'perioada': {'de_la': min(i['data'] for i in investigatii if i['data']),
                         'pana_la': max(i['data'] for i in investigatii if i['data'])},
            'generat': datetime.date.today().isoformat(),
        },
        'investigatii': investigatii,
        'puncte_proiectate': kml,
        'meteo': meteo,
    }
    os.makedirs(a.out, exist_ok=True)
    with open(os.path.join(a.out, 'data.json'), 'w', encoding='utf-8') as fh:
        json.dump(out, fh, ensure_ascii=False, indent=1)
    print('scris', os.path.join(a.out, 'data.json'), len(investigatii), 'investigații')


if __name__ == '__main__':
    main()
